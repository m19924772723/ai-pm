# -*- coding: utf-8 -*-
"""生成《AI产品经理工具栈与信息源》xlsx
   GitHub star 数 = 2026-09-22 实时查询；网站可访问性 = 同日 curl 实测
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(HERE, 'plans/AI产品经理工具栈与信息源.xlsx')

H_FILL = PatternFill('solid', fgColor='1F4E79'); H_FONT = Font(bold=True, color='FFFFFF', size=11)
THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
F = lambda c: PatternFill('solid', fgColor=c)
PRI_FILL = {'必装': F('C6EFCE'), '推荐': F('DEEAF6'), '可选': F('F2F2F2'), '进阶': F('FFF2CC')}
OK_FILL = {'200 正常': F('C6EFCE'), '403 防爬(站活着)': F('FFF2CC'), '000 本网络不可达(需代理)': F('FCE4D6')}

def head(ws, cols, widths):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.fill = H_FILL; cell.font = H_FONT; cell.alignment = CENTER; cell.border = BORDER
    for i, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.row_dimensions[1].height = 30

def body(ws, rows, h=48, center_cols=(), fill_map=None, fill_col=None):
    r = 2
    for row in rows:
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v); c.border = BORDER
            c.alignment = CENTER if i in center_cols else WRAP
            if fill_map and fill_col and i == fill_col:
                c.fill = fill_map.get(str(v), F('FFFFFF'))
        ws.row_dimensions[r].height = h
        r += 1
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(row)), r - 1)
    return r

wb = Workbook()

# ==========================================================
# Sheet 1: 必装软件
# ==========================================================
ws = wb.active; ws.title = '必装软件'
COLS = ['类别', '软件 / 工具', '是什么', '为什么 AI PM 要装', '开源地址（★）', '优先级', '备注']
W = [20, 26, 30, 40, 32, 9, 30]
ROWS = [
 # A 干活主力
 ('A 干活主力', 'Claude Code / Codex CLI', '命令行 AI 编程代理', '2026 AI PM 的标配：不用写代码也能让 AI 帮你把东西做出来', '—（已在用）', '必装', '你已在用，继续用'),
 ('A 干活主力', 'Cursor', 'AI 代码编辑器', '看代码、改原型、理解别人的开源项目', '—（闭源）', '必装', '读 GitHub 项目源码用它最快'),
 ('A 干活主力', 'dyad', '本地开源 AI 建站器（v0/Lovable/Bolt 替代）', '一句话生成完整网页应用，且完全本地跑', 'dyad-sh/dyad ★21.6k', '推荐', '想省 API 费用时用它'),
 ('A 干活主力', 'bolt.new', '一句话生成全栈网页', '快速验证想法，10 分钟出一个能点的原型', 'stackblitz/bolt.new ★16.6k', '推荐', '在线版，适合演示'),
 ('A 干活主力', 'OpenCode', '开源 AI 编程代理', '免费替代品，可接任意模型', '—（本机已装 skill）', '可选', ''),
 ('A 干活主力', 'we0', '面向 PM 的 AI 代码编辑器', '专门给产品经理做的，同类是 v0/bolt', 'we0-dev/we0 ★925', '可选', '可参考它的交互设计'),

 # B 本地模型
 ('B 本地模型', 'Ollama', '本地跑开源模型', '免费试各种模型、做本地 demo，不用花 API 钱', 'ollama/ollama ★181k', '必装', '一条命令跑通一个模型'),
 ('B 本地模型', 'Open WebUI', '本地版 ChatGPT 界面', '统一管理多模型，做产品 demo 的门面', 'open-webui/open-webui ★152k', '必装', '支持 Ollama + 任意 API'),
 ('B 本地模型', 'AnythingLLM', '本地知识库 + Agent 一体', '不用写代码就能体验完整 RAG 产品长什么样', 'Mintplex-Labs/anything-llm ★66k', '推荐', '做 L2 前先玩这个建立直觉'),
 ('B 本地模型', 'LM Studio', '图形化本地模型客户端', '不想碰命令行时的选择', '—（闭源免费）', '推荐', ''),
 ('B 本地模型', 'Jan', '离线 ChatGPT 替代', '100% 离线，隐私场景', 'janhq/jan ★44.6k', '可选', ''),

 # C RAG
 ('C RAG 知识库', 'RAGFlow', '开源 RAG 引擎', '快速搭一个能用的知识库——做 L2 项目的捷径', 'infiniflow/ragflow ★91k', '必装', '★ 做 L2 首选'),
 ('C RAG 知识库', 'LlamaIndex', '文档处理与 RAG 框架', '自己写 RAG 时的主力框架', 'run-llama/llama_index ★52k', '推荐', '文档最全'),
 ('C RAG 知识库', 'Langchain-Chatchat', '中文 RAG 应用', '中文场景的最佳参考实现', 'chatchat-space/Langchain-Chatchat ★38.7k', '推荐', '中文文档处理看它'),
 ('C RAG 知识库', 'Haystack', '生产级 RAG 编排框架', '从 demo 到生产的进阶', 'deepset-ai/haystack ★26.6k', '进阶', ''),
 ('C RAG 知识库', 'Quivr', '开箱即用 RAG 应用', '当参考实现读，别从零写', 'The-Vibe-Company/quivr ★39.5k', '可选', ''),

 # D Agent
 ('D Agent 框架', 'LangGraph', 'Agent 编排框架', 'L3 项目的地基，业界标准', '—（LangChain 生态）', '必装', '★ 做 L3 必学'),
 ('D Agent 框架', 'browser-use', '浏览器操作 Agent', 'L3 首选项目的核心库', 'browser-use/browser-use ★115k', '必装', '★ 2026 最热方向'),
 ('D Agent 框架', 'CrewAI', '角色化多 Agent 编排', '把任务拆给多个专职 Agent', 'crewAIInc/crewAI ★58.9k', '推荐', ''),
 ('D Agent 框架', 'AutoGen', '微软多 Agent 框架', '企业级多 Agent 方案', 'microsoft/autogen ★61.1k', '可选', ''),
 ('D Agent 框架', 'MetaGPT', '多 Agent 软件公司', '看多 Agent 怎么协作的经典案例', 'FoundationAgents/MetaGPT ★70.5k', '可选', ''),

 # E 评测观测
 ('E 评测与观测', 'Langfuse', '开源 Agent 评测与可观测平台', '★ 你项目要的"实测数字"靠它——追踪每次调用、算指标', 'langfuse/langfuse ★34.9k', '必装', '★ 最容易忽略、最能拉开差距'),
 ('E 评测与观测', 'Promptfoo', '提示词/Agent 测试与红队', '建评测集、跑对比实验、找漏洞', 'promptfoo/promptfoo ★25.4k', '必装', '★ 做 L1 评测就用它'),
 ('E 评测与观测', 'LiteLLM', '100+ 模型统一网关', '一个接口切所有模型，省钱且方便做模型对比', 'BerriAI/litellm ★59.4k', '推荐', '多模型对比实验必备'),
 ('E 评测与观测', 'Opik', 'LLM 评测监控', 'Langfuse 的替代方案', 'comet-ml/opik ★22.2k', '可选', ''),

 # F 数据
 ('F 数据与分析', 'DuckDB', '本地分析数据库', '练 SQL 不用装数据库，一个文件搞定', 'duckdb/duckdb ★41.6k', '推荐', '学 SQL 最轻的路'),
 ('F 数据与分析', 'Metabase', '开源 BI 看板', '做作品集案例③（数据驱动决策）时的看板', 'metabase/metabase ★49.4k', '推荐', ''),
 ('F 数据与分析', 'Jupyter', '数据分析环境', '数据处理与实验记录', '—（标准工具）', '推荐', ''),
 ('F 数据与分析', 'data-formulator', '微软 AI 数据分析工具', '自然语言→数据→图表，参考它的产品设计', 'microsoft/data-formulator ★17.2k', '可选', ''),

 # G MCP
 ('G MCP 生态', 'MCP 官方服务器集', '让 AI 接入你的工具（文件/数据库/浏览器）', '理解 Agent 怎么"动手做事"的关键', 'modelcontextprotocol/servers ★90.5k', '必装', '★ Agent 的能力来源'),
 ('G MCP 生态', 'awesome-mcp-servers', 'MCP 服务器大全', '找现成能力，不重复造轮子', 'punkpeye/awesome-mcp-servers ★95.4k', '推荐', ''),
 ('G MCP 生态', 'Playwright MCP', '浏览器自动化 MCP', '让 Agent 操作网页', 'microsoft/playwright-mcp ★37.5k', '可选', ''),

 # H 抓取
 ('H 抓取与自动化', 'Firecrawl', '网页数据 API', '给 Agent 喂干净网页数据', 'firecrawl/firecrawl ★183k', '推荐', '有免费额度'),
 ('H 抓取与自动化', 'Crawl4AI', '开源 LLM 友好爬虫', '自建抓取管道', 'unclecode/crawl4ai ★84k', '推荐', ''),
 ('H 抓取与自动化', 'Scrapling', '自适应爬虫框架', '反爬场景', 'D4Vinci/Scrapling ★83k', '可选', '本机已装对应 skill'),

 # I 设计
 ('I 设计原型', 'Figma', '交互原型设计', '作品集必备（前面已列入启动包）', '—（闭源）', '必装', 'AI 插件可加速'),
 ('I 设计原型', 'ai-design-skills', '设计 AI 产品的技能集', '参考别人怎么做 AI 产品设计的', 'Owl-Listener/ai-design-skills ★173', '可选', ''),

 # J PM 专用
 ('J PM 专用', 'Product-Manager-Skills', 'PM 技能库（77 个）', '★ 本机已装，你的主力工作流', 'deanpeters/Product-Manager-Skills ★7k', '已装', ''),
 ('J PM 专用', 'product-manager-prompts', 'PM 生成式 AI 提示词库', '同一作者的提示词合集，可补技能库', 'deanpeters/product-manager-prompts ★1.1k', '推荐', '★ 建议一并装'),
 ('J PM 专用', 'lenny-skills', 'Lenny 播客 86 个 PM 技能', 'Lenny\'s Newsletter 的 PM 技能化版本', 'RefoundAI/lenny-skills ★1.3k', '推荐', '★ 质量很高'),
 ('J PM 专用', 'pm-skills', '68 个 PM 技能（Triple Diamond）', '另一套成熟 PM 技能库，可对比借鉴', 'product-on-purpose/pm-skills ★691', '可选', ''),
 ('J PM 专用', 'AIPM-Wiki', 'AI 产品经理入门知识库', '★ 中文！面试题库 + AI 基础 + 案例拆解 + 学习路线', 'archlizheng/AIPM-Wiki ★320', '推荐', '★ 与你的目标完全对口'),
 ('J PM 专用', 'prd-writer', 'PRD 写作技能', '写 PRD 的专用技能，可对比本机 prd-development', 'chituai/prd-writer ★267', '可选', ''),
]
head(ws, COLS, W)
end = body(ws, ROWS, h=46, center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='共 %d 项。★ 数为 2026-09-22 实时查询 GitHub API。' % len(ROWS)).font = Font(bold=True, color='C00000')
ws.cell(row=end + 2, column=2, value='最优先装这 5 个：Claude Code/Cursor · Ollama · Open WebUI · Langfuse · Promptfoo').font = Font(bold=True)

# ==========================================================
# Sheet 2: GitHub 学习项目
# ==========================================================
ws = wb.create_sheet('GitHub学习项目')
COLS2 = ['学习目标', '项目', '★', '学什么', '怎么用（建议节奏）', '优先级']
W2 = [20, 36, 10, 40, 40, 9]
ROWS2 = [
 ('★ 抄 100+ 应用', 'Shubhamsaboo/awesome-llm-apps', '139.4k',
  '100+ 个可运行的 AI Agent / RAG 应用源码',
  '做 L1/L2 时直接找同类应用看它怎么实现——这是最省时间的一个仓库', '必学'),
 ('生成式 AI 入门', 'microsoft/generative-ai-for-beginners', '120.3k',
  '21 课，从零讲生成式 AI 应用开发',
  '通勤/睡前看，2 周过一遍，不求全懂', '必学'),
 ('Agent 入门', 'microsoft/ai-agents-for-beginners', '75.4k',
  '18 课，Agent 的原理与实现',
  '第 8 周（做 L3 前）看完', '必学'),
 ('提示词与上下文工程', 'dair-ai/Prompt-Engineering-Guide', '78.5k',
  '提示词、上下文工程、RAG、AI Agent 的系统指南',
  '第 1 周起当手册，随用随查', '必学'),
 ('中文 LLM 入门', 'datawhalechina/llm-cookbook', '24.7k',
  '吴恩达大模型系列课程中文版',
  '英文吃力时优先看这个', '推荐'),
 ('Agent 官方课程', 'huggingface/agents-course', '32.7k',
  'Hugging Face 官方 Agent 课程',
  '系统学 Agent 时用', '推荐'),
 ('中文 LLM 资源', 'WangRongsheng/awesome-LLM-resources', '9.0k',
  '多模态、Agent、辅助编程、MCP 等中文资料汇总',
  '找中文资料的第一站', '推荐'),
 ('中文模型/生态', 'AiHubCN/Awesome-Chinese-LLM', '22.8k',
  '中文大模型、垂直微调、数据集与教程',
  '做中文场景时查', '推荐'),
 ('Agent 项目清单', 'e2b-dev/awesome-ai-agents', '30.1k',
  '自主 Agent 项目大全',
  '找 L3/L4 项目灵感', '推荐'),
 ('工程化落地', 'liguodongiot/llm-action', '25.1k',
  '大模型工程化与应用落地的实战经验',
  '做 L4 平台项目时看', '推荐'),
 ('★ 学产品设计', 'asgeirtj/system_prompts_leaks', '68.0k',
  '各大 AI 产品的真实系统提示词',
  '★ 看顶级 AI 产品怎么设计提示词与边界——这是 AI PM 独有的学习材料', '必学'),
 ('★ 学产品设计', 'x1xhlol/system-prompts-and-models-of-ai-tools', '143.8k',
  'Cursor/Claude Code/Devin/Manus 等工具的系统提示词与内部工具',
  '同上，更全；做产品时对照它怎么约束模型行为', '必学'),
 ('中文 PM 知识库', 'archlizheng/AIPM-Wiki', '320',
  'AI 产品经理入门：面试题库 · AI 基础 · 案例拆解 · 学习路线',
  '★ 与你的目标完全对口，中文，先通读一遍', '必学'),
 ('读代码库', 'The-Pocket/PocketFlow-Tutorial-Codebase-Knowledge', '12.7k',
  '把代码库变成教程的方法',
  '读开源项目源码时用它当方法', '可选'),
 ('评测与可观测', 'langfuse/langfuse', '34.9k',
  '怎么给 LLM 应用做追踪与评测',
  '做 L1 评测时边用边学', '推荐'),
]
head(ws, COLS2, W2)
end = body(ws, ROWS2, h=54, center_cols=(1, 3, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ 最该先看两个：awesome-llm-apps（找参考实现）+ AIPM-Wiki（中文入门）').font = Font(bold=True, color='C00000')

# ==========================================================
# Sheet 3: 关注网站
# ==========================================================
ws = wb.create_sheet('关注网站')
COLS3 = ['语言', '类别', '网站', '网址', '看什么', '建议频率', '可访问性（实测）']
W3 = [7, 18, 24, 32, 40, 14, 22]
ROWS3 = [
 ('中文', '产品方法论', '人人都是产品经理', 'https://www.woshipm.com', '产品方法论、案例分析、面试经验——你的主信息源', '每周 2 篇精读', '200 正常'),
 ('中文', '商业科技', '36氪', 'https://36kr.com', 'AI 行业动态、公司战略、商业模式', '每天扫标题', '200 正常'),
 ('中文', 'AI 资讯', '机器之心', 'https://www.jiqizhixin.com', 'AI 技术进展与产业落地', '每周 3–5 篇', '200 正常'),
 ('中文', 'AI 资讯', '量子位', 'https://www.qbitai.com', 'AI 产品与行业新闻，中文里最快', '每天扫标题', '200 正常'),
 ('中文', '技术资讯', 'InfoQ', 'https://www.infoq.cn', '技术架构与工程实践，理解研发视角', '每周 2 篇', '200 正常'),
 ('中文', '工具与效率', '少数派', 'https://sspai.com', '工具评测、效率方法、产品体验分析', '每周 2 篇', '200 正常'),
 ('中文', '社区', '即刻', 'https://web.okjike.com', 'AI 从业者实时讨论，能拿到一手风向', '每天 10 分钟', '200 正常'),
 ('中文', '技术社区', '掘金', 'https://juejin.cn', 'AI 应用开发实战文章', '按需搜索', '200 正常'),
 ('中文', '问答', '知乎', 'https://www.zhihu.com', 'AI 产品讨论、面试经验、行业分析', '按需搜索', '403 防爬(站活着)'),
 ('中文', '求职', '牛客网', 'https://www.nowcoder.com', '★ 大厂真题、面经、内推——面试题库主要来源', '每周 3 次', '200 正常'),
 ('中文', '求职', '应届生求职网', 'https://www.yingjiesheng.com', '校招信息与进度', '每周 2 次', '200 正常'),
 ('中文', '求职', 'BOSS 直聘', 'https://www.zhipin.com', '实习岗位投递主渠道', '投递期每天', '200 正常'),
 ('英文', '★ PM 最前沿', "Lenny's Newsletter", 'https://www.lennysnewsletter.com', '★ 全球最好的产品经理 Newsletter，AI PM 内容极多', '每周 1 篇精读', '200 正常'),
 ('英文', '产品发布', 'Product Hunt', 'https://www.producthunt.com', '每天新产品，看 AI 产品形态风向', '每天扫标题', '403 防爬(站活着)'),
 ('英文', '技术社区', 'Hacker News', 'https://news.ycombinator.com', '技术圈风向，AI 讨论质量高', '每天扫标题', '000 本网络不可达(需代理)'),
 ('英文', 'AI 周报', 'The Batch (DeepLearning.AI)', 'https://www.deeplearning.ai/the-batch/', '吴恩达团队的 AI 周报，判断准确', '每周 1 期', '200 正常'),
 ('英文', 'AI 深度', 'Latent Space', 'https://www.latent.space', 'AI 工程与产品深度访谈', '每两周 1 篇', '200 正常'),
 ('英文', 'LLM 实践', "Simon Willison's Blog", 'https://simonwillison.net', 'LLM 实践第一手观察，更新极快', '每周 2 篇', '200 正常'),
 ('英文', '模型社区', 'Hugging Face', 'https://huggingface.co', '模型、数据集、demo（Spaces）', '按需', '000 本网络不可达(需代理)'),
 ('英文', '模型对比', 'Artificial Analysis', 'https://artificialanalysis.ai', '模型能力/价格/速度横向对比——选型必看', '做选型时', '200 正常'),
 ('英文', '模型竞技场', 'LMArena', 'https://lmarena.ai', '模型盲测排名', '做选型时', '000 本网络不可达(需代理)'),
 ('英文', '模型网关', 'OpenRouter', 'https://openrouter.ai', '一个接口调所有模型，看价格与可用性', '做选型时', '200 正常'),
 ('英文', '论文', 'arXiv', 'https://arxiv.org', '前沿论文（你写小论文也用）', '按需', '200 正常'),
 ('英文', '商业分析', 'Stratechery', 'https://stratechery.com', '科技商业战略分析，训练商业判断', '每周 1 篇', '200 正常'),
]
head(ws, COLS3, W3)
end = body(ws, ROWS3, h=42, center_cols=(1, 6, 7), fill_map=OK_FILL, fill_col=7)
ws.cell(row=end + 1, column=3, value='可访问性为 2026-09-22 本机 curl 实测：403=防爬但站活着；000=本网络不通，需代理。').font = Font(bold=True, color='C00000')
ws.cell(row=end + 2, column=3, value='每天固定 20 分钟：woshipm 精读 + 机器之心/量子位扫标题。不要刷太多，读深比读多重要。').font = Font(bold=True)

# ==========================================================
# Sheet 4: 上手顺序
# ==========================================================
ws = wb.create_sheet('上手顺序')
COLS4 = ['步骤', '什么时候', '装什么 / 关注什么', '为什么这个顺序']
W4 = [8, 18, 56, 46]
ROWS4 = [
 ('1', '本周（第 1 周）', 'Claude Code/Cursor（已在用）· Ollama · Open WebUI · woshipm + 机器之心',
  '先把"能跑模型"和"有信息源"这两件基础设施搭起来'),
 ('2', '第 3 周（做 L1 时）', 'Promptfoo（建评测集）· Langfuse（追踪）· AnythingLLM（体验完整 RAG 产品）',
  'L1 的核心是"评测方法"，先有度量工具再动手'),
 ('3', '第 5–8 周（做 L2 时）', 'RAGFlow 或 LlamaIndex · Langchain-Chatchat（中文参考）· 牛客网（开始建题库）',
  '做 RAG 项目时才需要 RAG 框架，早装只会吃灰'),
 ('4', '第 9–13 周（做 L3 时）', 'LangGraph · browser-use · Firecrawl · CrewAI',
  'L3 是 Agent 编排，这四样是地基'),
 ('5', '第 14 周起（包装）', 'Figma（原型）· Metabase（数据看板）· LiteLLM（多模型对比）',
  '作品集阶段要出图、出数据、出对比'),
 ('6', '持续（每天 20 分钟）', 'woshipm 精读 1 篇 · 量子位/机器之心扫标题 · Lenny\'s 每周 1 篇',
  '信息摄入要固定成习惯，而不是想起来才看'),
 ('7', '研一下起', 'MCP 服务器集 · awesome-mcp-servers · Hugging Face · LMArena',
  '做 L4 平台项目时，MCP 和模型选型会成为核心工作'),
]
head(ws, COLS4, W4)
end = body(ws, ROWS4, h=48, center_cols=(1, 2))
ws.cell(row=end + 1, column=3, value='★ 核心原则：用什么装什么，不要一次装完——装了不用的工具只会占硬盘和注意力。').font = Font(bold=True, color='C00000')
ws.cell(row=end + 2, column=3, value='★ 最优先的 5 个：Claude Code/Cursor · Ollama · Open WebUI · Langfuse · Promptfoo').font = Font(bold=True)

wb.save(OUT)
print('saved:', OUT)
print('sheets:', wb.sheetnames)
print('必装软件:', len(ROWS), '| 学习项目:', len(ROWS2), '| 网站:', len(ROWS3), '| 上手步骤:', len(ROWS4))
