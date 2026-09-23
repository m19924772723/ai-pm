# -*- coding: utf-8 -*-
"""加 sheet：立即启动（不依赖导师）+ 暂缓（等导师）"""
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, 'plans/待做产品项目清单.xlsx')

H_FILL = PatternFill('solid', fgColor='1F4E79'); H_FONT = Font(bold=True, color='FFFFFF', size=11)
THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
F = lambda c: PatternFill('solid', fgColor=c)
GRP_FILL = {'① 现在做': F('C6EFCE'), '② 等导师确定后': F('FCE4D6')}

wb = load_workbook(XLSX)

COLS = ['分组', '编号', '类别', '具体做什么', '为什么重要', '为什么现在能独立做', '耗时', '时间窗', '完成']
W = [15, 7, 12, 42, 34, 34, 9, 14, 7]

ROWS = [
 # ---------------- ① 现在做 ----------------
 ('① 现在做', 'A1', '学业线',
  '确定导师（第一标准：放不放实习）',
  '★ 唯一门槛动作——它不放实习，后面全链条作废',
  '不需要导师配合，是你主动去问、去选',
  '1–3 周', '第 1–3 周', None),

 ('① 现在做', 'B1', '项目线',
  '精读俞军《产品方法论》第 1 部分；记牢"用户价值 = 新体验 − 旧体验 − 替换成本"',
  '判断"值不值得做"的第一把尺子；写 PRD、做取舍、答面试题都要用',
  '纯自学，零依赖',
  '约 6 小时', '第 1 周', None),

 ('① 现在做', 'B2', '项目线',
  'Figma 装好，学会画框线原型',
  '原型是作品集的入场券，也是面试"你怎么表达方案"的载体',
  '纯自学，零依赖',
  '2–3 天', '第 2 周', None),

 ('① 现在做', 'B3', '项目线',
  '跑一遍 prd-development，摸清 PRD 的结构',
  '后面写自己的 PRD 时不用从零想框架',
  '技能库已装好，直接可用',
  '半天', '第 2 周', None),

 ('① 现在做', 'B4', '项目线',
  '拆解 1 个自己天天用的 AI 产品（用户 / JTBD / 竞品 / PRD 结构）',
  '练"结构化拆产品"的手感，同时为作品集案例①打底',
  '拆的是公开产品，零依赖',
  '1–2 天', '第 2 周', None),

 ('① 现在做', 'C1', '简历线',
  '建简历骨架文件（一页纸，先搭模块，不填空话）',
  '简历是随项目长出来的，先有容器才能往里装',
  '零依赖',
  '1 小时', '第 1 周', None),

 ('① 现在做', 'C2', '简历线',
  '写"教育背景"+"技能"两段（含 AI 工具链 / Python / Figma）',
  '这两段不依赖任何经历，现在就能定稿',
  '零依赖',
  '2 小时', '第 2 周', None),

 ('① 现在做', 'D1', '面试线',
  '建题库：收集 20 道大厂产品面试真题',
  '先把题型分布看清，才知道每种要练多少',
  '公开面经，零依赖',
  '3 小时', '第 1 周', None),

 ('① 现在做', 'D2', '面试线',
  '产品设计题 3 道（用六段式框架，口述并录音）',
  '★ 唯一有效练法是开口+录音；越早开始越不吃力',
  '不需要项目素材也能练（用公开产品当题目）',
  '3–4 小时', '第 2 周', None),

 ('① 现在做', 'D3', '面试线',
  '费米估算 3 道（四步拆解法）',
  '与项目完全无关的独立题型，适合现在打底',
  '纯方法论，零依赖',
  '3 小时', '第 3 周', None),

 ('① 现在做', 'E1', '资产与工具线',
  '注册 woshipm，每周精读 2 篇深度分析并写笔记',
  '建立行业感觉，且是长期复利动作',
  '零依赖，随时可做',
  '每周 2 小时', '第 1 周起持续', None),

 ('① 现在做', 'B20', '项目线',
  '★ 主项目启动：选 P3 文献阅读与对比 Agent 做选题定界',
  '不依赖导师/数据集，是你 100% 能独立完成的候选；先写出范围界定一页纸',
  'P3 不依赖任何外部条件（P1/P2 要等导师确认数据集）',
  '1 周', '第 3 周起', None),

 # ---------------- ② 等导师 ----------------
 ('② 等导师确定后', 'A2', '学业线', '小论文选题定下（图像缺陷检测方向优先）',
  '学术成果硬门槛', '需要导师给方向与判断', '3–4 周', '选定导师后', None),
 ('② 等导师确定后', 'A3', '学业线', '找导师确认：缺陷检测方向的数据集是否可行',
  '决定主项目走 P1/P2 还是 P3', '必须导师配合才有数据', '1 周', '选定导师后', None),
 ('② 等导师确定后', 'A4', '学业线', '小论文投出',
  '研一下必须完成的硬节点', '依赖选题与导师指导', '数月', '研一下 2027.2–7', None),
 ('② 等导师确定后', 'A5', '学业线', '大论文开题（选题须与实习兼容）',
  '毕业硬门槛', '依赖导师与小论文进度', '—', '研二上', None),
 ('② 等导师确定后', 'P1', '项目线', '【候选】包装稿合规预检助手',
  '北印行业契合度最高', '需行业资源与方向确认', '3–4 周', '导师确定后评估', None),
 ('② 等导师确定后', 'P2', '项目线', '【候选】印刷品缺陷检测辅助工具',
  '与小论文同向，一石二鸟', '需导师方向 + 数据集支持', '4–6 周', '导师确定后评估', None),
 ('② 等导师确定后', 'B16', '项目线', '作品集案例①：AI 产品复刻/深度拆解 + 改进方案',
  '作品集三案例之一', '建议先有项目经验再拆，深度才够', '2–3 周', '研一下', None),
]

if '立即启动' in wb.sheetnames:
    del wb['立即启动']
ws = wb.create_sheet('立即启动', 0)

# 标题行
ws.cell(row=1, column=1, value='分组').fill = H_FILL
for i, c in enumerate(COLS, start=1):
    cell = ws.cell(row=1, column=i, value=c)
    cell.fill = H_FILL; cell.font = H_FONT; cell.alignment = CENTER; cell.border = BORDER
for i, wd in enumerate(W, start=1):
    ws.column_dimensions[get_column_letter(i)].width = wd
ws.row_dimensions[1].height = 30

r = 2
for row in ROWS:
    for i, v in enumerate(row, start=1):
        if i == 9:  # 完成列留空
            v = None
        c = ws.cell(row=r, column=i, value=v); c.border = BORDER
        c.alignment = CENTER if i in (1, 2, 3, 7, 8, 9) else WRAP
        if i == 1 and v: c.fill = GRP_FILL.get(v, F('FFFFFF'))
    ws.row_dimensions[r].height = 46
    r += 1
ws.freeze_panes = 'D2'
ws.auto_filter.ref = 'A1:I%d' % (r - 1)

# 说明
ws.cell(row=r + 1, column=1, value='结论').font = Font(bold=True, size=12, color='C00000')
ws.cell(row=r + 1, column=4, value='现在只做「① 现在做」这 12 项——全部不依赖导师，且都是简单但重要的打底动作。').font = Font(bold=True, size=12, color='C00000')
ws.cell(row=r + 2, column=4, value='其中 A1 定导师是唯一门槛：它在第 1–3 周完成，完成后再解锁「② 等导师确定后」那 7 项。').font = Font(bold=True)
ws.cell(row=r + 3, column=4, value='主项目先选 P3（文献对比 Agent），因为它不依赖外部条件；导师确定后可再评估是否转 P1/P2。').font = Font(bold=True)
ws.cell(row=r + 4, column=4, value='前 3 周的产出很轻：1 份产品拆解 + 1 份 PRD 骨架 + 简历骨架 + 20 道题库 + 3 份设计题录音。').font = Font(bold=True)

wb.save(XLSX)
print('saved:', XLSX)
print('sheets:', wb.sheetnames)
now = [x for x in ROWS if x[0] == '① 现在做']
later = [x for x in ROWS if x[0] == '② 等导师确定后']
print('现在做:', len(now), '项 | 等导师:', len(later), '项')
