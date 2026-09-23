# -*- coding: utf-8 -*-
"""追加 7 个 sheet 到《AI产品经理工具栈与信息源.xlsx》：
   全栈开发 / 部署上线 / 测试与评测 / 运维与可观测 / 安全与合规 / 生命周期全景 / 学习资源
"""
import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, 'plans/AI产品经理工具栈与信息源.xlsx')

H_FILL = PatternFill('solid', fgColor='1F4E79'); H_FONT = Font(bold=True, color='FFFFFF', size=11)
THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
F = lambda c: PatternFill('solid', fgColor=c)
PRI_FILL = {'必装': F('C6EFCE'), '必读': F('C6EFCE'), '推荐': F('DEEAF6'), '可选': F('F2F2F2'), '进阶': F('FFF2CC')}
OK_FILL = {'200 正常': F('C6EFCE'), '301 跳转': F('FFF2CC'), '000 本网络不可达(需代理)': F('FCE4D6')}

wb = load_workbook(XLSX)

def mk(title, idx, cols, widths, rows, h=46, center_cols=(1, 5, 6), fill_map=None, fill_col=None):
    if title in wb.sheetnames:
        del wb[title]
    ws = wb.create_sheet(title, idx)
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.fill = H_FILL; cell.font = H_FONT; cell.alignment = CENTER; cell.border = BORDER
    for i, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.row_dimensions[1].height = 30
    r = 2
    for row in rows:
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v); c.border = BORDER
            c.alignment = CENTER if i in center_cols else WRAP
            if fill_map and fill_col and i == fill_col:
                c.fill = fill_map.get(str(v), F('FFFFFF'))
        ws.row_dimensions[r].height = h
        r += 1
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), r - 1)
    return ws, r

TC = ['环节', '工具 / 项目', '是什么', '为什么需要', '开源地址（★）', '优先级', '备注']
TW = [16, 28, 30, 38, 32, 9, 30]

# ============ 1. 全栈开发 ============
ws, end = mk('全栈开发', 4, TC, TW, [
 ('前端框架', 'Next.js', 'React 全栈框架', '主流 AI 产品的前端标准，Vercel 部署无缝', '—（开源）', '必装', '会改就行，不必精通'),
 ('UI 组件', 'shadcn/ui', '可复制的 UI 组件库', 'AI 产品最常用的界面组件，出活快', '—（开源）', '推荐', ''),
 ('后端 API', 'FastAPI', '高性能 Python API 框架', 'AI 服务的后端首选（模型调用都走它）', 'fastapi/fastapi ★102.5k', '必装', '★ 你的主力后端'),
 ('全栈模板', 'full-stack-fastapi-template', '开箱全栈模板', '省掉从零搭项目结构的时间', 'fastapi/full-stack-fastapi-template ★45.7k', '推荐', 'FastAPI+React+Postgres'),
 ('后端规范', 'fastapi-best-practices', '后端最佳实践合集', '避免写出只能跑 demo 的代码', 'zhanymkanov/fastapi-best-practices ★18.1k', '可选', ''),
 ('数据库', 'PostgreSQL + pgvector', '关系库 + 向量检索一体', '一个库同时存业务数据和向量，省事', '—（开源）', '必装', '★ 别再单独装向量库'),
 ('后端即服务', 'Supabase', '开箱后端（库+鉴权+存储）', '没有后端经验也能快速上线', '—（开源核心）', '推荐', '作品集项目提速神器'),
 ('快速界面', 'Streamlit / Gradio', '数据与 AI 应用快速出界面', '2 周做出可交互 demo 的最快路径', '—（开源）', '必装', '★ L1 项目就用它'),
 ('AI 聊天组件', 'Vercel AI SDK / assistant-ui', '流式聊天界面组件', '不用从零写对话 UI', '—（开源）', '推荐', ''),
 ('可视化编排', 'Dify / Langflow', '可视化搭 Agent 与工作流', '先看懂"一个 AI 应用由哪些环节组成"', 'langgenius/dify ★156.8k / langflow ★155.1k', '必装', '★ 做 L2 前先玩透'),
 ('自动化编排', 'n8n', '可视化工作流自动化（原生 AI）', '把 AI 接进业务流程，不用写代码', 'n8n-io/n8n ★205.7k', '推荐', ''),
 ('数据模型', 'SQLModel', 'Python 数据模型 ORM', 'FastAPI 生态配套', 'fastapi/sqlmodel ★18.3k', '可选', ''),
 ('命令行工具', 'Typer', '写 CLI 的框架', '给自己做自动化小工具', 'fastapi/typer ★20.0k', '可选', ''),
], center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ AI 全栈最小组合：FastAPI（后端）+ Streamlit（界面）+ PostgreSQL/pgvector（数据）+ Dify（编排）').font = Font(bold=True, color='C00000')

# ============ 2. 部署上线 ============
ws, end = mk('部署上线', 5, TC, TW, [
 ('容器化', 'Docker + docker-compose', '打包与运行环境', '保证"我这里能跑"= "别人那里也能跑"', '—（开源）', '必装', '★ 上线第一步'),
 ('镜像优化', 'dive', '查看镜像每一层', '把镜像从 GB 级压到百 MB 级', 'wagoodman/dive ★54.6k', '可选', ''),
 ('生产编排', 'Kubernetes', '容器编排平台', '大规模生产部署的事实标准', '—（开源）', '进阶', '先懂概念，暂不必上手'),
 ('GPU 集群', 'GPUStack', 'GPU 集群管理（vLLM/SGLang）', '自部署模型时的算力调度', 'gpustack/gpustack ★5.7k', '进阶', ''),
 ('模型服务', 'vLLM', '高吞吐 LLM 推理服务引擎', '自己部署模型对外提供 API 时的标准', 'vllm-project/vllm ★92.4k', '推荐', '有 GPU 时再上'),
 ('模型服务', 'SGLang', '高性能推理服务框架', 'vLLM 的替代/补充', 'sgl-project/sglang ★36.3k', '可选', ''),
 ('本地云', 'LocalStack', '本地模拟 AWS', '不花钱测部署流程', 'localstack/localstack ★65.1k', '可选', ''),
 ('CI/CD', 'GitHub Actions', '自动化测试与部署', '每次提交自动跑测试、自动部署', '—（平台）', '必装', '★ 作品集项目必须配'),
 ('CI/CD 安全', 'zizmor', 'GitHub Actions 静态扫描', '防止 CI 配置被利用', 'zizmorcore/zizmor ★6.6k', '可选', ''),
 ('Agent 生产模板', 'agent-starter-pack', '带 CI/CD 的 Agent 生产模板', '看 Google 怎么做生产级 Agent', 'GoogleCloudPlatform/agent-starter-pack ★6.6k', '推荐', ''),
 ('托管平台', 'Vercel / Railway / Fly.io', '一键部署托管', '不折腾服务器就能上线给人用', '—（平台）', '推荐', '★ 作品集上线用这个'),
 ('边缘部署', 'Cloudflare Workers', '边缘函数部署', '低延迟、免运维', '—（平台）', '可选', ''),
 ('AI 网关', 'LiteLLM', '100+ 模型统一网关', '一个接口切所有模型，做限流与成本控制', 'BerriAI/litellm ★59.4k', '必装', '★ 已在"必装软件"列出'),
 ('AI 网关', 'Portkey Gateway', '带护栏的 AI 网关', '路由 1600+ 模型 + 50+ 护栏', 'Portkey-AI/gateway ★13.1k', '推荐', ''),
 ('AI 网关', 'Bifrost / Plano', '企业级 AI 网关', '负载均衡、可观测、Agent 数据面', 'maximhq/bifrost ★8.2k / katanemo/plano ★7.1k', '可选', ''),
], center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ 上线最小闭环：Docker 打包 → GitHub Actions 自动测 → Vercel/Railway 部署 → LiteLLM 统一模型出口').font = Font(bold=True, color='C00000')

# ============ 3. 测试与评测 ============
ws, end = mk('测试与评测', 6, TC, TW, [
 ('评测框架', 'Promptfoo', '提示词/Agent/RAG 测试与红队', '建评测集、跑对比实验、找漏洞', 'promptfoo/promptfoo ★25.4k', '必装', '★ 已在"必装软件"列出'),
 ('评测框架', 'OpenAI Evals', '官方评测框架 + 基准库', '看业界标准怎么定义评测', 'openai/evals ★19.5k', '推荐', ''),
 ('★ PM 专属评测技能', 'evals-skills', 'AI Evals 技能集（配套"AI Evals for Engineers & PMs"课程）', '★ 专为 PM 写的评测技能，与你的岗位完全对口', 'hamelsmu/evals-skills ★1.7k', '必装', '★ 强烈建议装进技能库'),
 ('RAG 评测', 'Ragas', 'RAG 专用评测指标', '检索质量、答案忠实度等专业指标', '—（开源）', '推荐', '做 L2 时用'),
 ('LLM 单元测试', 'DeepEval', '像写单测一样测 LLM', '把 LLM 测试接进 CI', '—（开源）', '推荐', ''),
 ('Agent 红队', 'agentic_security', 'Agent 漏洞扫描 / AI 红队工具包', '主动找自己 Agent 的安全漏洞', 'msoedov/agentic_security ★2.0k', '推荐', '★ 安全测评加分项'),
 ('渗透测试', 'Strix', '开源 AI 渗透测试工具', '自动找应用漏洞并修复', 'usestrix/strix ★64.1k', '可选', ''),
 ('端到端测试', 'Playwright', '浏览器 E2E 测试', '保证每次改动不破坏已有功能', 'microsoft/playwright ★96.5k', '必装', '★ 前端测试标准'),
 ('Python 测试', 'pytest', 'Python 测试框架', '后端逻辑回归测试', '—（标准）', '必装', ''),
 ('组件测试', 'Storybook', 'UI 组件开发与测试', '前端组件隔离测试', 'storybookjs/storybook ★91.1k', '可选', ''),
 ('可观测评测', 'Langfuse', '追踪 + 评测一体', '把线上真实调用变成评测样本', 'langfuse/langfuse ★34.9k', '必装', '★ 已在"必装软件"列出'),
], center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ 评测是 AI PM 最容易被忽略、也最能拉开差距的能力：别人交"我做了个 Agent"，你交"覆盖率 62%→81%，失败集中在 X"').font = Font(bold=True, color='C00000')

# ============ 4. 运维与可观测 ============
ws, end = mk('运维与可观测', 7, TC, TW, [
 ('全链路追踪', 'Langfuse', 'LLM 调用追踪与评测', '每一次调用的输入输出、耗时、成本都可查', 'langfuse/langfuse ★34.9k', '必装', '★ 已在"必装软件"列出'),
 ('全链路追踪', 'Arize Phoenix', 'LLM 追踪与评测平台', '开源可自托管，看 trace 定位问题', '—（开源）', '推荐', ''),
 ('全链路追踪', 'Opik', 'LLM 评测监控', 'Langfuse 替代方案', 'comet-ml/opik ★22.2k', '可选', ''),
 ('成本与延迟', 'Helicone', 'LLM 成本/延迟/缓存监控', '控制 API 花销，做缓存省钱', '—（开源）', '推荐', '★ 省钱必备'),
 ('基础设施监控', 'Prometheus + Grafana', '指标采集与看板', '服务健康度、QPS、错误率', '—（开源）', '推荐', '运维标准组合'),
 ('实验跟踪', 'Weights & Biases', '实验与模型跟踪', '记录每次实验参数与结果', '—（平台）', '推荐', '做模型对比时用'),
 ('ML 平台', 'MLflow', 'ML 全流程平台', '模型注册、版本、部署管理', 'mlflow/mlflow ★28.1k', '推荐', ''),
 ('Token 优化', 'headroom', '压缩工具输出/日志/RAG 片段', '减少 20% token，直接省钱', 'headroomlabs-ai/headroom ★73.5k', '可选', ''),
 ('Token 优化', 'rtk', 'CLI 代理，减少 60–90% token', '常用开发命令的 token 消耗优化', 'rtk-ai/rtk ★81.4k', '可选', ''),
 ('告警与值班', 'Grafana Alerting / PagerDuty', '告警与值班', '出事能第一时间知道', '—（平台）', '可选', ''),
], center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ 运维最小闭环：Langfuse 看每次调用 → Helicone 控成本 → Grafana 看服务健康度').font = Font(bold=True, color='C00000')

# ============ 5. 安全与合规 ============
ws, end = mk('安全与合规', 8, TC, TW, [
 ('★ 标准必读', 'OWASP LLM Top 10', 'LLM 应用十大安全风险清单', '★ 做 AI 产品必须知道有哪些坑（提示词注入、数据泄露…）', 'owasp.org', '必读', '★ 花 1 小时通读'),
 ('★ 标准', 'NIST AI RMF', 'AI 风险管理框架', '企业级 AI 合规的通用语言', 'nist.gov', '必读', '面试聊合规时用'),
 ('输出护栏', 'Guardrails AI', '给 LLM 输出加约束与校验', '防止模型输出不合规内容/格式错乱', 'guardrails-ai/guardrails ★7.4k', '必装', '★ 生产必备'),
 ('输出护栏', 'NeMo Guardrails', '英伟达可编程护栏', '企业级对话安全控制', 'NVIDIA-NeMo/Guardrails ★7.2k', '推荐', ''),
 ('输入防护', 'LLM Guard', 'LLM 交互安全工具包', '检测注入、敏感信息、有害输入', 'protectai/llm-guard ★3.2k', '必装', '★ 输入侧第一道防线'),
 ('注入防护', 'Superagent', '防提示词注入/数据泄露', '给 Agent 加一层防护', 'superagent-ai/superagent ★6.8k', '推荐', ''),
 ('注入检测', 'Rebuff', '提示词注入检测器', '专测注入攻击', 'protectai/rebuff ★1.5k', '可选', ''),
 ('Agent 防火墙', 'Pipelock', 'MCP/Agent 出口安全防火墙', '控制 Agent 能访问什么', 'luckyPipewrench/pipelock ★897', '可选', ''),
 ('技能安全', 'SkillSpector', 'AI 技能安全扫描（NVIDIA）', '★ 装任何 AI 技能前先扫一遍', 'NVIDIA/SkillSpector ★18.1k', '推荐', '★ 本机已装技能库，建议扫一次'),
 ('安全评估', 'PurpleLlama', 'Meta 官方 LLM 安全工具集', '评估并改进 LLM 安全性', 'meta-llama/PurpleLlama ★4.4k', '推荐', ''),
 ('隐私合规', 'Microsoft Presidio', 'PII 识别与脱敏', '处理用户数据前先脱敏（个保法要求）', '—（开源）', '推荐', '★ 国内合规刚需'),
 ('密钥管理', 'OpenBao / Varlock / fnox', '密钥存储与分发', '绝不把 API Key 写进代码', 'openbao ★7.5k / varlock ★4.6k / fnox ★2.2k', '必装', '★ 基本工程素养'),
 ('安全资料', 'awesome-llm-security', 'LLM 安全资源清单', '系统了解这个领域', 'corca-ai/awesome-llm-security ★1.7k', '推荐', ''),
 ('安全资料', 'Awesome-LM-SSP', '大模型安全/隐私论文清单', '深入看研究', 'CryptoAILab/Awesome-LM-SSP ★2.1k', '可选', ''),
 ('攻防演示', 'LLM Security (greshake)', '应用集成 LLM 的攻击方式', '理解"为什么提示词注入是真问题"', 'greshake/llm-security ★2.1k', '推荐', ''),
], center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ 安全最小闭环：LLM Guard（输入）→ Guardrails AI（输出）→ Presidio（隐私）→ OpenBao（密钥）').font = Font(bold=True, color='C00000')

# ============ 6. 生命周期全景 ============
ws, end = mk('生命周期全景', 9,
 ['阶段', '要装什么（最小组合）', '练出的能力', '对应你的项目', '面试会问什么', '优先级'],
 [16, 44, 32, 18, 36, 9],
 [
 ('1 需求与设计', 'Figma · PM 技能库（已装）· AIPM-Wiki', '产品定义、取舍、PRD', '全部', '解决谁的问题？为什么这么做？', '必装'),
 ('2 全栈开发', 'FastAPI · Streamlit · PostgreSQL+pgvector · Supabase', '把想法变成能跑的产品', 'L1 / L2', '你负责哪部分？技术选型怎么定的？', '必装'),
 ('3 模型与 Agent', 'Ollama · vLLM · LangGraph · browser-use · Dify', 'Agent 编排、工具调用、失败处理', 'L2 / L3', '工作流为什么这么设计？失败怎么排查？', '必装'),
 ('4 测试与评测', 'Promptfoo · OpenAI Evals · DeepEval · Ragas · evals-skills', '可度量、可对比、可复现', 'L1 起全程', '效果怎么衡量？数字怎么来的？', '必装'),
 ('5 部署上线', 'Docker · GitHub Actions · Vercel/Railway · LiteLLM', '真的上线给真人用', 'L4', '有没有上线？谁在用？', '必装'),
 ('6 运维与可观测', 'Langfuse · Helicone · Grafana · Phoenix', '稳定运行、控成本、定位故障', 'L4', '线上出问题你怎么发现？', '推荐'),
 ('7 安全与合规', 'OWASP Top10 · LLM Guard · Guardrails AI · Presidio · OpenBao', '不出安全事故、满足合规', '全程', '提示词注入怎么防？用户数据怎么处理？', '必装'),
 ('8 数据与迭代', 'DuckDB · Metabase · AB 实验设计', '数据驱动决策', 'L4 / 案例③', '指标掉了怎么排查？怎么做实验？', '推荐'),
], h=52, center_cols=(1, 4, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='★ 这就是"AI 全栈产品经理"的完整能力地图：8 个阶段，每个阶段都有对应工具和面试问题。').font = Font(bold=True, color='C00000')
ws.cell(row=end + 2, column=2, value='★ 顺序原则：先用最小组合跑通一遍（L1），再逐阶段补深——不要一次把 8 个阶段全装齐。').font = Font(bold=True)

# ============ 7. 学习资源 ============
ws, end = mk('学习资源（运维部署安全）', 10,
 ['类别', '资源', '是什么', '学什么', '地址（★）', '优先级', '建议节奏'],
 [16, 30, 30, 34, 34, 9, 26],
 [
 ('MLOps 课程', 'mlops-zoomcamp', 'DataTalks 免费 MLOps 课程', '从实验到生产的完整 MLOps 流程', 'DataTalksClub/mlops-zoomcamp ★15.3k', '推荐', '做 L4 前过一遍'),
 ('LLMOps 实战', 'llm-twin-course', '端到端生产级 LLM/RAG 课程', 'LLMOps 最佳实践、部署、监控', 'decodingai-magazine/llm-twin-course ★4.4k', '推荐', '做 L3/L4 时跟做'),
 ('ML 工程', 'mlops-course (Made With ML)', '生产级 ML 应用设计', '设计→开发→部署→迭代全流程', 'GokuMohandas/mlops-course ★3.4k', '推荐', ''),
 ('AI 工程课', 'AI-Engineering-Lab', '24 周自学 AI 工程课程', 'Python → ML → LLM → RAG → 微调 → 部署', 'zorost/AI-Engineering-Lab ★308', '可选', '长期跟'),
 ('全栈 AI 课', 'ai-builders-curriculum', '全栈 AI 应用开发课程', '从零构建全栈 AI 应用', 'ai-builders-foundation/ai-builders-curriculum ★1.4k', '推荐', ''),
 ('DevOps 练习', 'devops-exercises', 'DevOps/SRE 题库（Linux/Docker/K8s/Prometheus）', '运维基本功', 'bregman-arie/devops-exercises ★84.6k', '可选', '按需查'),
 ('MLOps 清单', 'awesome-mlops / Awesome-LLMOps', 'MLOps/LLMOps 工具大全', '找工具、了解全景', 'visenger/awesome-mlops ★14.2k / tensorchord/Awesome-LLMOps ★5.9k', '推荐', ''),
 ('安全清单', 'awesome-llm-security', 'LLM 安全资源', '安全领域全景', 'corca-ai/awesome-llm-security ★1.7k', '推荐', ''),
 ('安全清单', 'awesome-ai-safety', 'AI 质量与安全论文/文章', '安全与质量前沿', 'Giskard-AI/awesome-ai-safety ★222', '可选', ''),
 ('标准', 'OWASP LLM Top 10', 'LLM 应用十大风险', '★ AI 产品安全的通用语言', 'owasp.org（200 正常）', '必读', '1 小时通读'),
 ('标准', 'NIST AI RMF', 'AI 风险管理框架', '企业合规框架', 'nist.gov（200 正常）', '必读', ''),
 ('平台文档', 'Vercel / Kubernetes / Docker 文档', '部署平台与容器文档', '部署与运维实操', 'vercel.com 200 · kubernetes.io 200 · docs.docker.com 000(需代理)', '推荐', '用到再查'),
 ('平台文档', 'W&B / Arize / LangSmith 文档', '实验跟踪与追踪文档', '可观测实操', 'wandb.ai 200 · arize.com 200 · docs.smith.langchain.com 200', '推荐', ''),
], h=46, center_cols=(1, 5, 6), fill_map=PRI_FILL, fill_col=6)
ws.cell(row=end + 1, column=2, value='可访问性为 2026-09-22 本机 curl 实测。000 = 本网络不通，需代理。').font = Font(bold=True, color='C00000')

wb.save(XLSX)
print('saved:', XLSX)
print('sheets:', wb.sheetnames)
